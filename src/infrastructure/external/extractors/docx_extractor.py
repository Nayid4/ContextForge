"""Word DOCX Document Text Extractor for Scoped RAG."""

from __future__ import annotations

import logging
from pathlib import Path

logger = logging.getLogger(__name__)


class DocxExtractor:
    """Extracts paragraphs and tables from Word (.docx) files using python-docx."""

    def extract(self, file_path: str) -> list[dict[str, str]]:
        path = Path(file_path)
        if not path.exists():
            return []

        chunks: list[dict[str, str]] = []
        try:
            from docx import Document

            doc = Document(str(path))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            for i, table in enumerate(doc.tables):
                paragraphs.append(f"[Tabla {i + 1}]")
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    paragraphs.append(" | ".join(cells))

            full_text = "\n\n".join(paragraphs).strip()
            if full_text:
                chunks.append({
                    "text": full_text,
                    "location": "Documento DOCX",
                })
        except Exception as exc:
            logger.warning("Error extrayendo texto de DOCX '%s': %s", file_path, exc)

        return chunks
