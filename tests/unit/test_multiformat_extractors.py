"""Unit tests for Multiformat Document Extractors (PDF, DOCX, XLSX, TXT)."""

from __future__ import annotations

from pathlib import Path
import pandas as pd
from docx import Document
from pypdf import PdfWriter
import pytest

from infrastructure.external.extractors import (
    DocxExtractor,
    ExcelExtractor,
    MultiFormatExtractor,
    PDFExtractor,
    TextExtractor,
)


def test_text_extractor(tmp_path: Path) -> None:
    extractor = TextExtractor()

    # Non-existent file
    assert extractor.extract(str(tmp_path / "missing.txt")) == []

    # Valid txt
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("Reglas técnicas del sistema.", encoding="utf-8")
    chunks = extractor.extract(str(txt_file))
    assert len(chunks) == 1
    assert "Reglas técnicas" in chunks[0]["text"]
    assert "notes.txt" in chunks[0]["location"]


def test_docx_extractor(tmp_path: Path) -> None:
    extractor = DocxExtractor()

    # Non-existent file
    assert extractor.extract(str(tmp_path / "missing.docx")) == []

    # Create dummy docx with paragraphs and tables
    doc_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Especificación de Seguridad", level=1)
    doc.add_paragraph("Párrafo explicativo del protocolo.")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Header A"
    table.rows[0].cells[1].text = "Header B"
    table.rows[1].cells[0].text = "Val 1"
    table.rows[1].cells[1].text = "Val 2"

    doc.save(str(doc_path))

    chunks = extractor.extract(str(doc_path))
    assert len(chunks) == 1
    assert "Especificación de Seguridad" in chunks[0]["text"]
    assert "Header A | Header B" in chunks[0]["text"]
    assert chunks[0]["location"] == "Documento DOCX"


def test_excel_extractor(tmp_path: Path) -> None:
    extractor = ExcelExtractor()

    # Non-existent file
    assert extractor.extract(str(tmp_path / "missing.xlsx")) == []

    # CSV extraction
    csv_path = tmp_path / "metrics.csv"
    df = pd.DataFrame({
        "Modulo": ["auth", "billing", "orders"],
        "Latencia_ms": [45, 120, 80],
    })
    df.to_csv(csv_path, index=False)

    csv_chunks = extractor.extract(str(csv_path))
    assert len(csv_chunks) >= 1
    assert "Modulo" in csv_chunks[0]["text"]
    assert "auth" in csv_chunks[0]["text"]
    assert "Filas 1-3" in csv_chunks[0]["location"]

    # Excel extraction
    xlsx_path = tmp_path / "report.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="KPIs", index=False)

    xlsx_chunks = extractor.extract(str(xlsx_path))
    assert len(xlsx_chunks) >= 1
    assert "Hoja 'KPIs'" in xlsx_chunks[0]["text"]
    assert "billing" in xlsx_chunks[0]["text"]


def test_pdf_extractor(tmp_path: Path) -> None:
    extractor = PDFExtractor()

    # Non-existent file
    assert extractor.extract(str(tmp_path / "missing.pdf")) == []

    # Generate a simple blank PDF page
    pdf_path = tmp_path / "dummy.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(pdf_path, "wb") as f:
        writer.write(f)

    # Empty blank page has no text, should return [] gracefully without errors
    chunks = extractor.extract(str(pdf_path))
    assert isinstance(chunks, list)


def test_multiformat_dispatcher(tmp_path: Path) -> None:
    dispatcher = MultiFormatExtractor()

    # Unsupported format returns empty list
    bin_file = tmp_path / "archive.zip"
    bin_file.write_bytes(b"PK000")
    assert dispatcher.extract(str(bin_file)) == []

    # Route to text
    txt_file = tmp_path / "query.sql"
    txt_file.write_text("SELECT * FROM projects;", encoding="utf-8")
    sql_chunks = dispatcher.extract(str(txt_file))
    assert len(sql_chunks) == 1
    assert "SELECT * FROM projects" in sql_chunks[0]["text"]

